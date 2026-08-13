from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "UniGather-新手使用操作手册.md"
OUTPUT = ROOT / "docs" / "UniGather-新手使用操作手册.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side in ("top", "bottom", "start", "end"):
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), "80" if side in ("top", "bottom") else "120")
                node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Microsoft YaHei", size=11, color="1F2937", bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("1F2937")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "UniGather 材料收集系统 · 新手使用操作手册"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=9, color="718096")
    footer = section.footer.paragraphs[0]
    footer.text = "Windows 单机版 · 版本 0.0.4"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run_font(run, size=9, color="718096")


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "E8F4FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color="1F4D78", bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    widths = [2700, 6660] if len(rows[0]) == 2 else [1800] * len(rows[0])
    if len(rows[0]) != 2:
        widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value.strip())
            set_run_font(r, size=9.5, color="1F2937", bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, "E8EEF5")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_markdown(doc, text):
    lines = text.splitlines()
    idx = 0
    in_code = False
    code_lines = []
    while idx < len(lines):
        line = lines[idx]
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.right_indent = Inches(0.2)
                p.paragraph_format.space_after = Pt(8)
                for code_line in code_lines:
                    r = p.add_run(code_line + "\n")
                    set_run_font(r, name="Consolas", size=9.5, color="334155")
                in_code = False
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue
        if not line.strip():
            idx += 1
            continue
        if line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.space_before = Pt(5)
            r = p.add_run("▰")
            set_run_font(r, size=7, color="25A7C8")
            idx += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            p = doc.add_paragraph(style=f"Heading {level}")
            r = p.add_run(title)
            set_run_font(r, size={1:16, 2:13, 3:12}[level], color={1:"2E74B5", 2:"2E74B5", 3:"1F4D78"}[level], bold=True)
            idx += 1
            continue
        if line.startswith("> "):
            add_callout(doc, line[2:].strip())
            idx += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while idx < len(lines) and lines[idx].startswith("|"):
                if not re.match(r"^\|\s*[-:| ]+\|\s*$", lines[idx]):
                    table_lines.append([part.strip() for part in lines[idx].strip().strip("|").split("|")])
                idx += 1
            if table_lines:
                add_table(doc, table_lines)
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        bullet = re.match(r"^-\s+(.*)$", line)
        if numbered or bullet:
            style = "List Number" if numbered else "List Bullet"
            content = (numbered or bullet).group(1)
            p = doc.add_paragraph(style=style)
            p.add_run(content)
            for run in p.runs:
                set_run_font(run, size=11)
            idx += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = False
        # Basic inline code emphasis for readability.
        parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", line)
        for part in parts:
            if not part:
                continue
            bold = part.startswith("**") and part.endswith("**")
            code = part.startswith("`") and part.endswith("`")
            value = part[2:-2] if bold else (part[1:-1] if code else part)
            r = p.add_run(value)
            set_run_font(r, name="Consolas" if code else "Microsoft YaHei", size=10.5 if code else 11, color="1F4D78" if code else "1F2937", bold=bold)
        idx += 1


def main():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("UniGather 材料收集系统")
    set_run_font(r, size=24, color="0B2545", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("新手使用操作手册 · 重点覆盖收集任务与独立邮件发送")
    set_run_font(r, size=12, color="718096")
    add_callout(doc, "第一次使用请按：邮箱配置 → 通讯录 → 材料归档 → 收集任务 / 邮件发送 的顺序操作。")

    source = SOURCE.read_text(encoding="utf-8")
    # Skip the markdown title and subtitle because the Word cover already contains them.
    source = re.sub(r"^# UniGather 材料收集系统\n\n## 新手使用操作手册.*?---\n", "", source, flags=re.S)
    parse_markdown(doc, source)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
